"""予定一覧 (1 枚目) と同じ形式の Excel 入力テンプレートを作る.

シート = 1 か月。列は  日 | メンバー1 | メンバー2 | ... | (無題: 科の予定)。
土日・祝日は日番号セルを赤で塗る (元ファイルと同じ)。
``parse_plan.parse_plan_xlsx`` でそのまま読める。
"""

from __future__ import annotations

import calendar
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .parse_plan import DEPT_KEY, Plan
from .skeleton import split_top_level

# 祝日 (2026-09 〜 2027-03). 必要になったら追記する.
JP_HOLIDAYS: dict[dt.date, str] = {
    dt.date(2026, 9, 21): "敬老の日",
    dt.date(2026, 9, 22): "国民の休日",
    dt.date(2026, 9, 23): "秋分の日",
    dt.date(2026, 10, 12): "スポーツの日",
    dt.date(2026, 11, 3): "文化の日",
    dt.date(2026, 11, 23): "勤労感謝の日",
    dt.date(2027, 1, 1): "元日",
    dt.date(2027, 1, 11): "成人の日",
    dt.date(2027, 2, 11): "建国記念の日",
    dt.date(2027, 2, 23): "天皇誕生日",
    dt.date(2027, 3, 21): "春分の日",
    dt.date(2027, 3, 22): "振替休日",
}

# 年末年始休暇 (12/29〜1/3). 大学病院の休診日に合わせて既定で赤にする.
YEAR_END_DAYS = {(12, 29), (12, 30), (12, 31), (1, 2), (1, 3)}

RED = "FF0000"
FONT_NAME = "MS PGothic"


def holiday_name(date: dt.date, year_end: bool = True) -> str | None:
    if date in JP_HOLIDAYS:
        return JP_HOLIDAYS[date]
    if year_end and (date.month, date.day) in YEAR_END_DAYS:
        return "年末年始"
    if date.weekday() == 5:
        return "土"
    if date.weekday() == 6:
        return "日"
    return None


def month_range(start: str, end: str) -> list[tuple[int, int]]:
    y, m = (int(v) for v in start.split("-"))
    ye, me = (int(v) for v in end.split("-"))
    out = []
    while (y, m) <= (ye, me):
        out.append((y, m))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def _style_sheet(ws, members: list[str], ndays: int, col_widths: list[float]) -> None:  # noqa: ANN001
    font = Font(name=FONT_NAME, size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hair = Side(style="thin", color="BFBFBF")
    black = Side(style="thin", color="000000")
    ncols = len(members) + 2  # 日 + メンバー + dept
    for ci, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    for r in range(1, ndays + 2):
        nlines = 1
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = font
            cell.alignment = center
            bottom = black if r == 1 else hair
            cell.border = Border(bottom=bottom)
            if isinstance(cell.value, str):
                # 折り返し行数の見積もり (全角 1 文字 ≒ 幅 2)
                width = col_widths[c - 1] if c - 1 < len(col_widths) else 10
                for ln in cell.value.split("\n"):
                    est = sum(2 if ord(ch) > 0x7F else 1 for ch in ln)
                    nlines = max(nlines, -(-est // max(int(width), 1)) + cell.value.count("\n"))
        ws.row_dimensions[r].height = 18 * nlines
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:{get_column_letter(ncols)}{ndays + 1}"


def add_month_sheet(wb: Workbook, year: int, month: int, members: list[str], plan: Plan | None = None, year_end: bool = True):  # noqa: ANN001
    ws = wb.create_sheet(f"{year:04d}-{month:02d}")
    ndays = calendar.monthrange(year, month)[1]
    # ヘッダ: A1 空, メンバー名, 最後は無題 (科の予定)
    for ci, name in enumerate(members, start=2):
        ws.cell(row=1, column=ci, value=name)
    dept_col = len(members) + 2
    for day in range(1, ndays + 1):
        r = day + 1
        c = ws.cell(row=r, column=1, value=day)
        date = dt.date(year, month, day)
        hol = holiday_name(date, year_end)
        if hol:
            c.fill = PatternFill(fill_type="solid", start_color=RED, end_color=RED)
        if plan is not None:
            row = plan.days.get(day) or {}
            for ci, name in enumerate(members, start=2):
                if row.get(name):
                    ws.cell(row=r, column=ci, value=row[name])
            if row.get(DEPT_KEY):
                # 科の予定は元ファイルと同じくセル内改行で項目を分ける (読み込み時に ", " に戻る)
                ws.cell(row=r, column=dept_col, value="\n".join(split_top_level(row[DEPT_KEY])))
    widths = [6] + [max(18, 44 - 12 * i) for i in range(len(members))] + [30]
    _style_sheet(ws, members, ndays, widths)
    return ws


def add_guide_sheet(wb: Workbook, members: list[str], months: list[tuple[int, int]], year_end: bool) -> None:
    ws = wb.create_sheet("使い方", 0)
    font = Font(name=FONT_NAME, size=11)
    bold = Font(name=FONT_NAME, size=11, bold=True)
    lines: list[tuple[str, Font]] = [
        ("予定一覧 入力シートの使い方", bold),
        ("", font),
        ("・シート名 = 年月 (YYYY-MM)。1 行目がメンバー名、A 列が日、最後の無題列が科全体の予定です。", font),
        ("・各メンバーの列に、その日の予定をカンマ区切りで書きます。書き方の例:", font),
        ("    外来 (AM/PM)            午前・午後とも外来", font),
        ("    外来 (PM)               午後のみ", font),
        ("    休暇                    終日休暇 (昼間)", font),
        ("    休暇, 平鹿 (PM)         午前休暇、午後は平鹿", font),
        ("    男鹿 (AM)               午前は男鹿へ出張", font),
        ("    男鹿前日                翌朝から出張のため夜の当番に入れない", font),
        ("    医療安全管理部担当者会議 (16時～)   時刻付きの会議", font),
        ("    web会議 (19時～)        夜の予定", font),
        ("    OSCE外部評価者 (佐賀)   場所付き (終日)", font),
        ("・A 列の日番号が赤い日 = 土日・祝日" + (" (12/29〜1/3 の年末年始も赤にしています)" if year_end else "") + "。休日を変えるときはセルの塗りつぶしを変えてください。", font),
        ("・記入例として 2026-09 シートに実際の予定を入れてあります。", font),
        ("・このファイルはそのまま  python -m duty_roster parse-plan <このファイル> --month YYYY-MM  で読み込めます。", font),
        ("", font),
        ("祝日 (このファイルで赤にした日):", bold),
    ]
    for (y, m) in months:
        names = []
        for day in range(1, calendar.monthrange(y, m)[1] + 1):
            d = dt.date(y, m, day)
            h = holiday_name(d, year_end)
            if h and h not in ("土", "日"):
                names.append(f"{m}/{day} {h}")
        lines.append((f"    {y}-{m:02d}: " + (", ".join(names) if names else "(祝日なし)"), font))
    for r, (text, f) in enumerate(lines, start=1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = f
        c.alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 110
    ws.sheet_view.showGridLines = False


def make_plan_template(
    path: str,
    start: str,
    end: str,
    members: list[str],
    example: Plan | None = None,
    year_end: bool = True,
) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    months = month_range(start, end)
    if example is not None:
        add_month_sheet(wb, example.year, example.month, example.members, example, year_end)
    for y, m in months:
        add_month_sheet(wb, y, m, members, None, year_end)
    all_months = ([(example.year, example.month)] if example else []) + months
    add_guide_sheet(wb, members, all_months, year_end)
    # 最初に開いたとき最初の入力月を表示
    first = wb[f"{months[0][0]:04d}-{months[0][1]:02d}"]
    wb.active = wb.sheetnames.index(first.title)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path
