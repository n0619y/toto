"""Roster → Excel (.xlsx) カレンダー生成. 2枚目 PDF と同じレイアウトを再現する.

レイアウト (列):  A = 行ラベル, B.. = 7 日 × 3 コマ = 21 列
レイアウト (行):  タイトル / 曜日 / 時刻 / [日付, メンバー×N] / 時刻 / [日付, メンバー×N] ...
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import (
    COLOR_FREE,
    COLOR_HOLIDAY,
    SLOT_TIMES,
    WEEKDAY_JA,
    Roster,
    build_cells,
)

try:  # openpyxl >= 3.1
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    HAS_RICH_TEXT = True
except ImportError:  # pragma: no cover
    HAS_RICH_TEXT = False


class Style:
    """見た目の調整値 (フォント・サイズ・幅・高さ)."""

    font_name = "MS PGothic"
    font_size = 8.0
    note_font_size = 5.5
    title_font_size = 11.0
    row_height = 15.75  # pt
    slot_col_width = 4.57  # Excel 幅単位 (文字数)
    label_col_width = 5.2
    thin = Side(style="thin", color="000000")
    thick = Side(style="medium", color="000000")


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=hex_color, end_color=hex_color)


def render_workbook(roster: Roster, style: Style | None = None) -> Workbook:
    st = style or Style()
    wb = Workbook()
    ws = wb.active
    ws.title = f"{roster.year}-{roster.month:02d}"

    members = roster.members
    nslots = 21
    first_col = 2  # B
    last_col = first_col + nslots - 1
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    base_font = Font(name=st.font_name, size=st.font_size)

    # 列幅
    ws.column_dimensions["A"].width = st.label_col_width
    for c in range(first_col, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = st.slot_col_width

    def put(row: int, col: int, value=None, fill: str | None = None, font: Font | None = None):  # noqa: ANN001
        cell = ws.cell(row=row, column=col)
        if value is not None:
            cell.value = value
        cell.alignment = center
        cell.font = font or base_font
        if fill:
            cell.fill = _fill(fill)
        return cell

    # ---- ヘッダ -----------------------------------------------------------------
    row = 1
    ws.row_dimensions[row].height = st.row_height
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    put(row, first_col, roster.title or roster.default_title(), font=Font(name=st.font_name, size=st.title_font_size))

    row = 2
    ws.row_dimensions[row].height = st.row_height
    put(row, 1)
    for di in range(7):
        c0 = first_col + di * 3
        ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c0 + 2)
        put(row, c0, WEEKDAY_JA[di])
        for k in (1, 2):
            put(row, c0 + k)
    weekday_row = row

    def time_row(r: int) -> None:
        ws.row_dimensions[r].height = st.row_height
        put(r, 1, "時刻")
        for di in range(7):
            for si in range(3):
                put(r, first_col + di * 3 + si, SLOT_TIMES[si])

    row = 3
    time_row(row)
    time_rows = [row]
    day_rows: list[int] = []
    member_rows: dict[tuple[int, str], int] = {}

    weeks = roster.weeks()
    for wi, week in enumerate(weeks):
        if wi > 0:
            row += 1
            time_row(row)
            time_rows.append(row)
        # 日付行
        row += 1
        ws.row_dimensions[row].height = st.row_height
        day_rows.append(row)
        put(row, 1)
        for di, day in enumerate(week):
            c0 = first_col + di * 3
            ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c0 + 2)
            holiday = (day is not None and roster.is_holiday(day)) or (day is None and di >= 5)
            fill = COLOR_HOLIDAY if holiday else COLOR_FREE
            put(row, c0, day, fill=fill)
            for k in (1, 2):
                put(row, c0 + k, fill=fill)
        # メンバー行
        for member in members:
            row += 1
            ws.row_dimensions[row].height = st.row_height
            member_rows[(wi, member)] = row
            put(row, 1, member)

    last_row = row

    # ---- セル (結合済み) ---------------------------------------------------------
    for cell in build_cells(roster):
        r = member_rows[(cell.week_index, cell.member)]
        c0 = first_col + cell.start
        c1 = first_col + cell.end
        if c1 > c0:
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        for c in range(c0, c1 + 1):
            put(r, c, fill=cell.color)
        top_left = ws.cell(row=r, column=c0)
        top_left.value = _cell_value(cell.lines, st)
        if cell.cross:
            for c in range(c0, c1 + 1):
                x = ws.cell(row=r, column=c)
                x.border = Border(diagonal=st.thin, diagonalUp=True, diagonalDown=True)

    # ---- 罫線 ----------------------------------------------------------------------
    # 太線: 外枠, ラベル列の右, 日付の境界, 時刻行の上下. それ以外は細線.
    thick_cols = {1, last_col} | {first_col + di * 3 - 1 for di in range(1, 7)}  # 右側が太線になる列
    thick_bottom_rows = {weekday_row - 1, last_row} | set(time_rows) | {r - 1 for r in time_rows}
    for r in range(weekday_row, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            existing = cell.border
            left = st.thick if c == 1 else st.thin
            right = st.thick if c in thick_cols else st.thin
            top = st.thick if (r - 1) in thick_bottom_rows else st.thin
            bottom = st.thick if r in thick_bottom_rows else st.thin
            cell.border = Border(
                left=left,
                right=right,
                top=top,
                bottom=bottom,
                diagonal=existing.diagonal,
                diagonalUp=existing.diagonalUp,
                diagonalDown=existing.diagonalDown,
            )

    # ---- 印刷設定 --------------------------------------------------------------------
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.sheet_view.showGridLines = False
    return wb


def _cell_value(lines: tuple[str, ...], st: Style):  # noqa: ANN001
    if not lines:
        return None
    if len(lines) == 1 or not HAS_RICH_TEXT:
        return "\n".join(lines)
    # 1 行目は通常サイズ, 2 行目以降 (補足) は小さく
    rich = CellRichText()
    rich.append(TextBlock(InlineFont(rFont=st.font_name, sz=st.font_size), lines[0]))
    for extra in lines[1:]:
        rich.append(TextBlock(InlineFont(rFont=st.font_name, sz=st.note_font_size), "\n" + extra))
    return rich


def save_xlsx(roster: Roster, path: str, style: Style | None = None) -> str:
    wb = render_workbook(roster, style)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


def xlsx_to_pdf(xlsx_path: str, pdf_path: str | None = None) -> str:
    """LibreOffice で xlsx → PDF 変換する (soffice が必要)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) が見つかりません。PDF 変換には LibreOffice が必要です。")
    xlsx_path = os.path.abspath(xlsx_path)
    if pdf_path is None:
        pdf_path = os.path.splitext(xlsx_path)[0] + ".pdf"
    pdf_path = os.path.abspath(pdf_path)
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [
                soffice,
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{tmp}/profile",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp,
                xlsx_path,
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=180,
        )
        produced = os.path.join(tmp, os.path.splitext(os.path.basename(xlsx_path))[0] + ".pdf")
        os.makedirs(os.path.dirname(pdf_path), exist_ok=True)
        shutil.move(produced, pdf_path)
    return pdf_path
