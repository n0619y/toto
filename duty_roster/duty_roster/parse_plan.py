"""1枚目 (メンバー予定一覧) の読み込み.

対応形式:

* PDF  … Excel の一覧表を印刷したもの (Schedule_202609.PDF の形式)
         列: 日 | メンバー1 | メンバー2 | ... | (無題: 科の予定)   赤塗りの日番号 = 休日
* xlsx … 同じ列構成のシート (1 行目がヘッダ、A 列が日)
* yaml … 本ツールの plan.yaml

plan.yaml の形::

    month: 2026-09
    members: [豊野, 仲本, 佐々木]
    holidays: [5, 6, 12, 13, 19, 20, 21, 22, 23, 26, 27]
    days:
      1:
        豊野: 外来 (AM/PM), web会議 (19時～)
        dept: 一般定期健康診断
      2:
        ...
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import pymupdf
import yaml

from .model import dump_yaml

DEPT_KEY = "dept"  # 科全体の予定 (メンバー列ではない無題列)


@dataclass
class Plan:
    year: int
    month: int
    members: list[str]
    holidays: set[int] = field(default_factory=set)
    # days[day][member or "dept"] = "外来 (AM/PM), web会議 (19時～)"
    days: dict[int, dict[str, str]] = field(default_factory=dict)

    def entry(self, day: int, member: str) -> str:
        return (self.days.get(day) or {}).get(member, "") or ""

    def to_dict(self) -> dict:
        return {
            "month": f"{self.year:04d}-{self.month:02d}",
            "members": list(self.members),
            "holidays": sorted(self.holidays),
            "days": {d: {k: v for k, v in row.items() if v} for d, row in sorted(self.days.items())},
        }

    @classmethod
    def from_dict(cls, data: dict) -> "Plan":
        y, m = (int(v) for v in str(data["month"]).split("-"))
        plan = cls(y, m, list(data["members"]), set(int(d) for d in data.get("holidays", [])))
        for day, row in (data.get("days") or {}).items():
            plan.days[int(day)] = {str(k): str(v) for k, v in (row or {}).items() if v}
        return plan

    def save(self, path: str) -> None:
        with open(path, "w", encoding="utf-8") as f:
            f.write(dump_yaml(self.to_dict()))

    @classmethod
    def load(cls, path: str) -> "Plan":
        with open(path, encoding="utf-8") as f:
            return cls.from_dict(yaml.safe_load(f))


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def parse_plan_pdf(path: str, year: int, month: int) -> Plan:
    """予定一覧 PDF を読む.

    ページは 90° 回転して印刷されているため、PyMuPDF の未回転座標では
    「x = 行方向 (下に増える)」「y = 列方向 (左に行くほど大きい)」になる。
    回転行列で表示座標に直してから処理する。
    """
    doc = pymupdf.open(path)
    page = doc[0]
    mat = page.rotation_matrix  # 未回転座標 → 表示座標

    spans: list[tuple[pymupdf.Rect, str]] = []
    for block in page.get_text("dict")["blocks"]:
        for line in block["lines"]:
            for sp in line["spans"]:
                if sp["text"].strip():
                    spans.append((pymupdf.Rect(sp["bbox"]) * mat, sp["text"].strip()))
    spans = [(r.normalize(), t) for r, t in spans]

    # 赤塗り = 休日
    red_rects = []
    for dr in page.get_drawings():
        fill = dr.get("fill")
        if fill and round(fill[0], 1) == 1.0 and round(fill[1], 1) == 0.0 and round(fill[2], 1) == 0.0:
            red_rects.append((dr["rect"] * mat).normalize())

    # ヘッダ行 (メンバー名) = 最上段の文字列
    top_y = min(r.y0 for r, _ in spans)
    header = sorted(((r, t) for r, t in spans if r.y0 < top_y + 5), key=lambda s: s[0].x0)
    members = [t for _, t in header if not t.isdigit()]

    # 日番号 (最左列の数字)
    numbers = [(r, int(t)) for r, t in spans if t.isdigit() and r.y0 > top_y + 5]
    left_x = min(r.x0 for r, _ in numbers)
    day_spans = sorted(((r, n) for r, n in numbers if r.x0 < left_x + 20), key=lambda s: s[0].y0)
    days = [n for _, n in day_spans]
    if days != list(range(1, len(days) + 1)):
        raise ValueError(f"日番号が連続していません: {days}")

    # 行境界 = 隣接する日番号の中点 (先頭はヘッダとの中点)
    centers = [(r.y0 + r.y1) / 2 for r, _ in day_spans]
    header_cy = sum((r.y0 + r.y1) / 2 for r, _ in header) / len(header)
    bounds = [(header_cy + centers[0]) / 2] + [(a + b) / 2 for a, b in zip(centers, centers[1:])] + [1e9]

    # 列境界 = ヘッダ中心同士の中点. 最後のメンバー列の右端は「ヘッダ中心に対して左右対称」とみなし、
    # それより右にある文字は無題の科予定列 (dept) とする.
    hx = [(r.x0 + r.x1) / 2 for r, t in header if not t.isdigit()]
    col_bounds = [(a + b) / 2 for a, b in zip(hx, hx[1:])]
    if len(hx) >= 2:
        col_bounds.append(hx[-1] + (hx[-1] - col_bounds[-1]))
    else:
        col_bounds.append(hx[-1] + 60)
    day_col_right = left_x + 25

    def column_of(cx: float) -> str | None:
        if cx < day_col_right:
            return None
        idx = sum(1 for b in col_bounds if cx > b)
        if idx < len(members):
            return members[idx]
        return DEPT_KEY

    plan = Plan(year, month, members)
    for r, n in day_spans:
        if any(rr.contains(pymupdf.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)) for rr in red_rects):
            plan.holidays.add(n)

    cell_lines: dict[tuple[int, str], list[tuple[pymupdf.Rect, str]]] = {}
    for r, t in spans:
        cy = (r.y0 + r.y1) / 2
        if cy < top_y + 5:
            continue
        row = next((i for i in range(len(days)) if bounds[i] <= cy < bounds[i + 1]), None)
        if row is None:
            continue
        col = column_of((r.x0 + r.x1) / 2)
        if col is None:
            continue
        cell_lines.setdefault((days[row], col), []).append((r, t))

    # 列ごとの最大文字幅 (折り返し判定用)
    col_max_w: dict[str, float] = {}
    for (d, col), items in cell_lines.items():
        for r, _ in items:
            col_max_w[col] = max(col_max_w.get(col, 0), r.width)

    for (d, col), items in cell_lines.items():
        items.sort(key=lambda s: s[0].y0)
        text = _join_lines([(r, t) for r, t in items], col_max_w.get(col, 0))
        plan.days.setdefault(d, {})[col] = text
    return plan


def _join_lines(items: list[tuple[pymupdf.Rect, str]], col_w: float) -> str:
    """セル内の複数行を 1 つの文字列にする.

    * 前の行が列幅いっぱい (= 自動折り返し) なら連結
    * 括弧が閉じていなければ連結
    * それ以外は別項目として ", " で区切る
    """
    out: list[str] = []
    for i, (r, t) in enumerate(items):
        if not out:
            out.append(t)
            continue
        prev_rect, prev_text = items[i - 1]
        unbalanced = prev_text.count("(") + prev_text.count("（") > prev_text.count(")") + prev_text.count("）")
        wrapped = col_w > 0 and prev_rect.width >= col_w * 0.9
        if unbalanced or wrapped:
            sep = " " if (prev_text[-1].isascii() and prev_text[-1].isalnum()) or (t[0].isascii() and t[0].isalnum()) else ""
            out[-1] = out[-1] + sep + t
        else:
            out.append(t)
    return ", ".join(out)


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------


def parse_plan_xlsx(path: str, year: int, month: int, sheet: str | None = None) -> Plan:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    members = [h for h in header[1:] if h]
    plan = Plan(year, month, members)
    red_days: set[int] = set()
    for row_cells in ws.iter_rows(min_row=2):
        first = row_cells[0]
        if first.value is None:
            continue
        try:
            day = int(first.value)
        except (TypeError, ValueError):
            continue
        fill = first.fill.start_color.rgb if first.fill and first.fill.fill_type == "solid" else None
        if isinstance(fill, str) and fill.upper().endswith("FF0000"):
            red_days.add(day)
        for ci, cell in enumerate(row_cells[1:], start=1):
            key = header[ci] if ci < len(header) and header[ci] else DEPT_KEY
            if cell.value not in (None, ""):
                plan.days.setdefault(day, {})[key] = re.sub(r"\s*\n\s*", ", ", str(cell.value).strip())
    plan.holidays = red_days
    return plan


def load_plan(path: str, year: int | None = None, month: int | None = None) -> Plan:
    low = path.lower()
    if low.endswith((".yaml", ".yml")):
        return Plan.load(path)
    if year is None or month is None:
        raise ValueError("PDF / xlsx を読むには --month YYYY-MM が必要です")
    if low.endswith(".pdf"):
        return parse_plan_pdf(path, year, month)
    if low.endswith((".xlsx", ".xlsm")):
        return parse_plan_xlsx(path, year, month)
    raise ValueError(f"未対応の形式: {path}")
