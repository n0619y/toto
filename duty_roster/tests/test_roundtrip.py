"""samples/ の 2 つの PDF を使った往復テスト.

* 元カレンダー PDF → roster.yaml → PDF 再描画 → 抽出 が元と一致すること
* 予定一覧 PDF が正しく読めること
* 予定文字列 → コマ変換のルール
"""

from __future__ import annotations

import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from duty_roster.check import analyze  # noqa: E402
from duty_roster.compare import compare_calendars, compare_roster_to_calendar, roster_from_calendar  # noqa: E402
from duty_roster.extract_calendar import extract_calendar  # noqa: E402
from duty_roster.model import Roster, SlotSpec, build_cells  # noqa: E402
from duty_roster.parse_plan import parse_plan_pdf  # noqa: E402
from duty_roster.render_pdf import render_pdf  # noqa: E402
from duty_roster.render_xlsx import save_xlsx  # noqa: E402
from duty_roster.skeleton import build_skeleton, events_to_slots, split_top_level  # noqa: E402

CAL = os.path.join(ROOT, "samples", "Sep_2026_1_calendar.pdf")
PLAN = os.path.join(ROOT, "samples", "Schedule_202609_plan.pdf")
ROSTER = os.path.join(ROOT, "data", "2026-09", "roster.yaml")


def test_extract_original_calendar():
    cal = extract_calendar(CAL)
    assert cal.members == ["豊野", "仲本", "佐々木"]
    assert cal.title == "Schedule for September 2026 (Cardiology)"
    assert [s.color for s in cal.slots[(1, "佐々木")]] == ["Y", "Y", "Y"]
    assert [s.label for s in cal.slots[(4, "豊野")]] == ["休暇", "佐賀", "佐賀"]
    assert cal.slots[(30, "豊野")][0].cross is True
    assert cal.slots[(9, "豊野")][2].lines == ("1st", "(19時～)")


def test_roster_yaml_matches_original():
    roster = Roster.load(ROSTER)
    assert compare_roster_to_calendar(roster, extract_calendar(CAL)) == []


def test_roster_from_calendar_roundtrip(tmp_path):
    cal = extract_calendar(CAL)
    roster = roster_from_calendar(cal, 2026, 9)
    assert roster.holidays == {21, 22, 23}
    path = tmp_path / "r.yaml"
    roster.save(str(path))
    assert compare_roster_to_calendar(Roster.load(str(path)), cal) == []


def test_render_pdf_reproduces_original(tmp_path):
    roster = Roster.load(ROSTER)
    pdf = render_pdf(roster, str(tmp_path / "out.pdf"))
    diffs = compare_calendars(extract_calendar(CAL), extract_calendar(pdf))
    assert diffs == []


def test_render_xlsx(tmp_path):
    from openpyxl import load_workbook

    roster = Roster.load(ROSTER)
    path = save_xlsx(roster, str(tmp_path / "out.xlsx"))
    ws = load_workbook(path).active
    assert ws["B1"].value == "Schedule for September 2026 (Cardiology)"
    assert ws["A3"].value == "時刻"
    # 1 日 (火) の佐々木 = 1st 終日 → E..G 結合, 黄色
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "E7:G7" in merged
    assert ws["E7"].fill.start_color.rgb.endswith("FFFF00")


def test_build_cells_merging():
    roster = Roster(2026, 9, ["A"])
    roster.days[1] = {"A": [SlotSpec.parse(s) for s in ("外来", "外来", "-")]}
    roster.days[2] = {"A": [SlotSpec.parse(s) for s in ("外来", "1st", "1st")]}
    cells = [c for c in build_cells(roster) if c.lines]
    # 1 日 AM+PM の外来 は結合. 2 日 AM の外来は 1 日夜 (空き) で切れる
    spans = {(c.start, c.end, c.text) for c in cells}
    assert (3, 4, "外来") in spans
    assert (6, 6, "外来") in spans
    assert (7, 8, "1st") in spans


def test_slotspec_syntax():
    assert SlotSpec.parse("外来/2nd") == SlotSpec(role="2nd", event="外来")
    assert SlotSpec.parse("1st+一般外来当番").label_lines() == ("1st", "一般外来当番")
    assert SlotSpec.parse("×").cross is True
    assert SlotSpec.parse("-").is_free
    assert SlotSpec.parse("NICU/1st").color(False) == "FFFF00"
    assert SlotSpec.parse("休暇").color(True) == "B1A0C7"
    assert SlotSpec.parse("-").color(True) == "92CDDC"


def test_parse_plan_pdf():
    plan = parse_plan_pdf(PLAN, 2026, 9)
    assert plan.members == ["豊野", "仲本", "佐々木"]
    assert plan.holidays == {5, 6, 12, 13, 19, 20, 21, 22, 23, 26, 27}
    assert plan.days[1]["豊野"] == "外来 (AM/PM), web会議 (19時～)"
    assert plan.days[1]["dept"] == "一般定期健康診断"
    assert plan.days[2]["dept"] == "Cath ? (心室中隔欠損, 乳児), 一般定期健康診断"
    assert plan.days[16]["仲本"] == "市立秋田 (PM), 男鹿前日"
    assert plan.days[24]["豊野"] == "外来 (AM/PM), 附属病院運営会議 (15時30分～)"
    assert plan.days[11]["dept"] == "運動負荷心筋核医学検査 1名"


@pytest.mark.parametrize(
    "text, expected",
    [
        ("外来 (AM/PM), web会議 (19時～)", ["外来", "外来", "Web会議"]),
        ("外来 (PM)", ["-", "外来", "-"]),
        ("外来 (AM), 休暇 (PM)", ["外来", "休暇", "-"]),
        ("休暇", ["休暇", "休暇", "-"]),
        ("休暇, 平鹿 (PM)", ["休暇", "平鹿", "-"]),
        ("男鹿 (AM)", ["男鹿", "-", "-"]),
        ("男鹿前日", ["-", "-", "-"]),
        ("医療安全管理部担当者会議 (16時～)", ["-", "会議", "-"]),
        ("休暇, 空港発 (15時50分)", ["休暇", "空港発", "空港発"]),
        ("空港着 (11時25分)", ["空港着", "-", "-"]),
        ("OSCE外部評価者 (佐賀)", ["佐賀", "佐賀", "佐賀"]),
        ("市立秋田 (PM), 男鹿前日", ["-", "市立秋田", "-"]),
    ],
)
def test_events_to_slots(text, expected):
    specs, _ = events_to_slots(text)
    assert [s.to_text() for s in specs] == expected


def test_split_top_level():
    assert split_top_level("Cath ? (心室中隔欠損, 乳児), 一般定期健康診断") == ["Cath ? (心室中隔欠損, 乳児)", "一般定期健康診断"]


def test_skeleton_matches_plan_events():
    plan = parse_plan_pdf(PLAN, 2026, 9)
    roster, comments = build_skeleton(plan)
    assert roster.holidays == {21, 22, 23}
    assert [s.to_text() for s in roster.slots(9, "豊野")] == ["休暇", "平鹿", "-"]
    assert "仲本: 男鹿前日  ※男鹿前日 (夜は当番不可)" in comments[9]


def test_check_original_has_one_first_per_slot():
    st = analyze(Roster.load(ROSTER))
    assert st.warnings == []
