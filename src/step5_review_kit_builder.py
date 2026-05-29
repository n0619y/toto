"""ステップ5: 監修チェックリスト自動添付。

toto先生が短時間で監修・修正指示を完了できるパッケージを生成します。
生成物（output/review/<テーマ名>_<日付>/ に格納）:
  (1) review_checklist.xlsx … 5シート構成の監修メインファイル
  (2) review_summary.md      … 人間がパッと読める要約版
  (3) review_package.html    … ブラウザで開ける一覧ビューア
"""

from __future__ import annotations

import html
import re
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .utils.citation_tracker import CitationTracker
from .utils.claude_client import ClaudeClient
from .utils.common import (
    DISCLAIMER_TEXT,
    OUTPUT_DIR,
    count_chars,
    ensure_dir,
    get_logger,
    load_review_checklist,
    load_settings,
    reviewer_display,
    reviewer_name,
    reviewer_title,
    slugify,
    today_iso,
    today_str,
)

logger = get_logger()

# 自動チェック用の正規表現・ワードリスト ------------------------------------
# 代表的な市販薬・薬剤の商品名（検出したら「一般名にすべき」と警告）
PRODUCT_NAME_PATTERN = re.compile(
    r"(カロナール|アンヒバ|ムコダイン|ムコソルバン|ホクナリン|メプチン|"
    r"アスベリン|ペリアクチン|ザイザル|アレグラ|アレジオン|クラリチン|"
    r"オノン|シングレア|タミフル|ゾフルーザ|イナビル|リレンザ|"
    r"ポンタール|ボルタレン|ロキソニン|セルテクト|ビオフェルミン|"
    r"ナウゼリン|プリンペラン|セレキノン|ラックビー)"
)

# 緊急症状ワード（冒頭に警告があるべきトピックの検出）
EMERGENCY_WORDS = [
    "けいれん", "痙攣", "意識障害", "意識がない", "呼吸困難", "呼吸が苦しい",
    "ぐったり", "脱水", "チアノーゼ", "3か月未満", "生後3か月",
    "反応が乏しい", "顔色が悪い", "唇が紫",
]


def _read(path: str) -> str:
    """ファイルを安全に読み込む（無ければ空文字）。"""
    p = Path(path)
    return p.read_text(encoding="utf-8") if p.exists() else ""


def run_auto_checks(article_md: str, youtube_md: str, shorts_md: str) -> Dict:
    """商品名・緊急ワード・免責文・監修者表記の自動チェックを行う。"""
    all_text = "\n".join([article_md, youtube_md, shorts_md])

    product_hits = sorted(set(PRODUCT_NAME_PATTERN.findall(all_text)))
    emergency_hits = sorted({w for w in EMERGENCY_WORDS if w in all_text})

    # 受診目安セクションの有無
    has_consult = bool(re.search(r"受診(の)?目安|救急|当日|翌日", all_text))
    # 免責文の有無
    has_disclaimer = DISCLAIMER_TEXT[:15] in all_text or "診断・治療は必ず医療機関" in all_text

    # 監修者表記の各成果物への挿入確認
    name = reviewer_name()
    reviewer_in = {
        "記事": name in article_md,
        "本編台本": name in youtube_md,
        "Shorts台本": name in shorts_md,
    }
    reviewer_all = all(reviewer_in.values())

    return {
        "product_hits": product_hits,
        "emergency_hits": emergency_hits,
        "has_consult": has_consult,
        "has_disclaimer": has_disclaimer,
        "reviewer_in": reviewer_in,
        "reviewer_all": reviewer_all,
    }


def _claude_risk_assessment(theme: str, article_md: str, checks: Dict) -> Dict:
    """Claudeにリスク判定と要注意箇所TOP5抽出を依頼する。失敗時はフォールバック。"""
    client = ClaudeClient()
    prompt = f"""あなたは小児科医の監修を補助するアシスタントです。
以下の記事ドラフトを医療安全の観点でチェックし、医師の確認が必須と思われる箇所を抽出してください。

# テーマ
{theme}

# 自動チェック結果
- 商品名検出: {checks['product_hits'] or 'なし'}
- 緊急症状ワード: {checks['emergency_hits'] or 'なし'}
- 受診目安セクション: {'あり' if checks['has_consult'] else 'なし'}
- 免責文: {'あり' if checks['has_disclaimer'] else 'なし'}

# 記事ドラフト（抜粋）
{article_md[:6000]}

# 出力（厳密に次のJSONのみ）
{{
  "risk_level": "高 or 中 or 低",
  "risk_reason": "リスク判定の理由を1〜2文で",
  "top5": [
    {{"quote": "医師確認が必要な記述を原文抜粋", "why": "理由"}}
  ]
}}
top5 は最大5件。"""
    try:
        import json

        raw = client.complete(prompt, max_tokens=2500, temperature=0.2).strip()
        if raw.startswith("```"):
            raw = raw.split("```", 2)[1]
        s, e = raw.find("{"), raw.rfind("}")
        return json.loads(raw[s : e + 1])
    except Exception as exc:
        logger.warning("Claudeリスク判定に失敗: %s — フォールバック判定を使用", exc)
        # 機械的フォールバック
        level = "高" if (checks["product_hits"] or not checks["has_disclaimer"]) else (
            "中" if checks["emergency_hits"] and not checks["has_consult"] else "低"
        )
        return {
            "risk_level": level,
            "risk_reason": "自動判定（Claude応答が得られませんでした）。",
            "top5": [],
        }


# === Excel 生成 =============================================================
HEADER_FILL = PatternFill("solid", fgColor="4FB3D9")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="2C6E8F")
WRAP = Alignment(wrap_text=True, vertical="top")
YELLOW = PatternFill("solid", fgColor="FFF3B0")
RED = PatternFill("solid", fgColor="F8B6B6")


def _style_header(ws, row, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = WRAP


def _autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _build_excel(
    path: Path, theme: str, citations: CitationTracker, checks: Dict,
    risk: Dict, products: Dict,
) -> None:
    """5シート構成の監修Excelを生成する。"""
    wb = Workbook()

    # --- Sheet1: サマリー ---
    ws = wb.active
    ws.title = "サマリー"
    ws["A1"] = "監修サマリー"
    ws["A1"].font = TITLE_FONT
    conf = citations.confidence_breakdown()
    org_cat = citations.org_category_breakdown()
    rows = [
        ("監修者", reviewer_display()),
        ("テーマ名", theme),
        ("作成日", today_iso()),
        ("想定読者・視聴者", load_settings()["content"]["target_audience"]),
        ("緊急度トピックの有無", "あり" if checks["emergency_hits"] else "なし"),
        ("リスク判定", risk.get("risk_level", "-")),
        ("", ""),
        ("記事 文字数", products.get("article_chars", 0)),
        ("本編台本 文字数", products.get("youtube_chars", 0)),
        ("Shorts台本 文字数", products.get("shorts_chars", 0)),
        ("本編スライド数", products.get("main_slide_count", 0)),
        ("Shortsスライド数", products.get("shorts_slide_count", 0)),
        ("", ""),
        ("出典: 学会", org_cat["学会"]),
        ("出典: 公的機関", org_cat["公的機関"]),
        ("出典: 海外ガイドライン", org_cat["海外ガイドライン"]),
        ("出典: その他", org_cat["その他"]),
        ("信頼度A / B / C", f"{conf['A']} / {conf['B']} / {conf['C']}"),
        ("", ""),
        ("監修完了日", "（ご記入ください）"),
        ("署名", "（ご記入ください）"),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = WRAP
    _autosize(ws, [24, 60])

    # --- Sheet2: 主張×出典 対応表 ---
    ws2 = wb.create_sheet("主張×出典")
    headers = ["主張ID", "主張文", "出典URL", "機関", "信頼度", "ガイドライン年度",
               "監修判定(OK/要修正/削除)", "コメント"]
    _style_header(ws2, 1, headers)
    for r, c in enumerate(citations.items, start=2):
        ws2.cell(row=r, column=1, value=c.claim_id)
        ws2.cell(row=r, column=2, value=c.claim).alignment = WRAP
        ws2.cell(row=r, column=3, value=c.url)
        ws2.cell(row=r, column=4, value=c.org)
        ws2.cell(row=r, column=5, value=c.confidence)
        ws2.cell(row=r, column=6, value=c.guideline_year or "不明")
        ws2.cell(row=r, column=7, value="")  # 監修判定（記入欄）
        ws2.cell(row=r, column=8, value="").alignment = WRAP
    _autosize(ws2, [10, 45, 40, 22, 8, 14, 22, 30])
    # 条件付き書式: 「要修正」=黄, 「削除」=赤
    from openpyxl.formatting.rule import CellIsRule

    last = max(ws2.max_row, 2)
    rng = f"G2:G{last}"
    ws2.conditional_formatting.add(
        rng, CellIsRule(operator="equal", formula=['"要修正"'], fill=YELLOW))
    ws2.conditional_formatting.add(
        rng, CellIsRule(operator="equal", formula=['"削除"'], fill=RED))

    # --- Sheet3: 必須チェック項目 ---
    ws3 = wb.create_sheet("必須チェック項目")
    _style_header(ws3, 1, ["カテゴリー", "ID", "チェック内容", "重要度", "チェック", "コメント"])
    r = 2
    for cat in load_review_checklist().get("categories", []):
        for item in cat.get("items", []):
            ws3.cell(row=r, column=1, value=cat["name"])
            ws3.cell(row=r, column=2, value=item["id"])
            ws3.cell(row=r, column=3, value=item["text"]).alignment = WRAP
            ws3.cell(row=r, column=4, value=item.get("priority", ""))
            ws3.cell(row=r, column=5, value="")  # チェック欄
            ws3.cell(row=r, column=6, value="").alignment = WRAP
            r += 1
    _autosize(ws3, [16, 10, 50, 10, 10, 30])

    # --- Sheet4: 成果物別レビュー ---
    ws4 = wb.create_sheet("成果物別レビュー")
    row = 1
    # 記事レビュー（見出しごと）
    ws4.cell(row=row, column=1, value="■ 記事レビュー").font = TITLE_FONT
    row += 1
    _style_header(ws4, row, ["見出し", "本文プレビュー", "OK/修正", "コメント"])
    row += 1
    for head, prev in products.get("article_sections", []):
        ws4.cell(row=row, column=1, value=head).alignment = WRAP
        ws4.cell(row=row, column=2, value=prev).alignment = WRAP
        ws4.cell(row=row, column=3, value="")
        ws4.cell(row=row, column=4, value="").alignment = WRAP
        row += 1
    row += 1
    # 本編動画レビュー（スライドごと）
    ws4.cell(row=row, column=1, value="■ 本編動画レビュー").font = TITLE_FONT
    row += 1
    _style_header(ws4, row, ["タイトル", "本文", "ナレーション", "OK/修正", "コメント"])
    row += 1
    for s in products.get("main_slides", []):
        ws4.cell(row=row, column=1, value=getattr(s, "title", "")).alignment = WRAP
        ws4.cell(row=row, column=2, value=getattr(s, "body", "")).alignment = WRAP
        ws4.cell(row=row, column=3, value=getattr(s, "narration", "")).alignment = WRAP
        row += 1
    row += 1
    # Shortsレビュー
    ws4.cell(row=row, column=1, value="■ Shortsレビュー").font = TITLE_FONT
    row += 1
    _style_header(ws4, row, ["タイトル", "本文", "ナレーション", "OK/修正", "コメント"])
    row += 1
    for s in products.get("shorts_slides", []):
        ws4.cell(row=row, column=1, value=getattr(s, "title", "")).alignment = WRAP
        ws4.cell(row=row, column=2, value=getattr(s, "body", "")).alignment = WRAP
        ws4.cell(row=row, column=3, value=getattr(s, "narration", "")).alignment = WRAP
        row += 1
    _autosize(ws4, [28, 45, 45, 10, 25])

    # --- Sheet5: 修正指示記入欄 ---
    ws5 = wb.create_sheet("修正指示記入欄")
    ws5["A1"] = "修正指示（フリーフォーマット）"
    ws5["A1"].font = TITLE_FONT
    ws5["A2"] = ("ここに修正指示を自由に記入してください。\n"
                 "次回 `python main.py --revise <このファイルのパス>` で反映の土台にできます。")
    ws5["A2"].alignment = WRAP
    _style_header(ws5, 4, ["対象成果物", "対象箇所(ID/見出し)", "修正内容", "優先度"])
    for r in range(5, 30):  # 記入用の空行
        for c in range(1, 5):
            ws5.cell(row=r, column=c, value="").alignment = WRAP
    _autosize(ws5, [20, 28, 60, 12])
    ws5.column_dimensions["A"].width = 20

    wb.save(str(path))


# === Markdown サマリー ======================================================
def _build_summary_md(path: Path, theme: str, checks: Dict, risk: Dict,
                      citations: CitationTracker) -> None:
    conf = citations.confidence_breakdown()
    org_cat = citations.org_category_breakdown()
    reviewer_lines = "\n".join(
        f"  - {k}: {'✅ あり' if v else '❌ なし'}" for k, v in checks["reviewer_in"].items()
    )
    top5 = risk.get("top5", [])
    top5_md = "\n".join(
        f"{i}. 「{html.unescape(t.get('quote',''))}」\n   → {t.get('why','')}"
        for i, t in enumerate(top5, 1)
    ) or "（特記事項なし、またはClaude未実行）"

    md = f"""# 監修依頼: {reviewer_display()}

- **テーマ**: {theme}
- **作成日**: {today_iso()}

## リスク判定: {risk.get('risk_level','-')}
{risk.get('risk_reason','')}

## 自動チェック結果
- **商品名検出**: {('⚠️ ' + ', '.join(f'`{p}`' for p in checks['product_hits'])) if checks['product_hits'] else '✅ 検出なし（一般名のみ）'}
- **緊急症状ワード**: {('⚠️ ' + ', '.join(checks['emergency_hits'])) if checks['emergency_hits'] else 'なし'}
- **受診目安セクション**: {'✅ あり' if checks['has_consult'] else '❌ なし（要追加）'}
- **免責文**: {'✅ あり' if checks['has_disclaimer'] else '❌ なし（要追加）'}
- **出典数**: {len(citations)} 件（信頼度 A:{conf['A']} / B:{conf['B']} / C:{conf['C']}）
- **出典機関の偏り**: 学会{org_cat['学会']} / 公的{org_cat['公的機関']} / 海外{org_cat['海外ガイドライン']} / その他{org_cat['その他']}
- **監修者表記「{reviewer_name()}」の全成果物への挿入確認**: {'✅ 全成果物OK' if checks['reviewer_all'] else '❌ 不足あり（下記）'}
{reviewer_lines}

## 要注意箇所 TOP5（医師確認推奨）
{top5_md}

---
> {DISCLAIMER_TEXT}
"""
    path.write_text(md, encoding="utf-8")


# === HTML ビューア ==========================================================
def _build_html_viewer(path: Path, theme: str, checks: Dict, risk: Dict,
                       article_md: str) -> None:
    design = load_settings()["design"]
    import markdown as md_lib

    article_html = md_lib.markdown(article_md, extensions=["extra", "tables"])
    checklist_html = ""
    for cat in load_review_checklist().get("categories", []):
        checklist_html += f"<h3>{html.escape(cat['name'])}</h3><ul>"
        for item in cat.get("items", []):
            checklist_html += (
                f"<li><input type='checkbox'> [{item['id']}] "
                f"{html.escape(item['text'])} "
                f"<span class='pri'>({item.get('priority','')})</span></li>"
            )
        checklist_html += "</ul>"

    auto_html = f"""
      <li>商品名検出: {'⚠️ ' + ', '.join(checks['product_hits']) if checks['product_hits'] else '✅ なし'}</li>
      <li>緊急ワード: {', '.join(checks['emergency_hits']) or 'なし'}</li>
      <li>受診目安: {'✅' if checks['has_consult'] else '❌'}</li>
      <li>免責文: {'✅' if checks['has_disclaimer'] else '❌'}</li>
      <li>監修者表記: {'✅ 全成果物' if checks['reviewer_all'] else '❌ 不足あり'}</li>
      <li>リスク判定: <b>{risk.get('risk_level','-')}</b></li>
    """

    doc = f"""<!DOCTYPE html>
<html lang="ja"><head><meta charset="utf-8">
<meta name="author" content="{reviewer_display()}">
<title>監修パッケージ - {html.escape(theme)}</title>
<style>
  body {{ font-family: "{design['font_main']}", sans-serif; margin: 0; color: #333; }}
  header {{ background: {design['primary_color']}; color: #fff; padding: 16px 24px; font-size: 1.2em; font-weight: bold; }}
  .wrap {{ display: flex; gap: 0; }}
  .left {{ width: 38%; padding: 20px; background: {design['secondary_color']}; height: 90vh; overflow-y: auto; }}
  .right {{ width: 62%; padding: 24px; height: 90vh; overflow-y: auto; }}
  h2 {{ color: {design['primary_color']}; }}
  h3 {{ color: {design['accent_color']}; }}
  .pri {{ color: #999; font-size: 0.85em; }}
  .auto {{ background: #fff; border-radius: 8px; padding: 10px 16px; }}
  li {{ margin: 6px 0; }}
</style></head>
<body>
<header>{reviewer_display()} ｜ 監修パッケージ: {html.escape(theme)}</header>
<div class="wrap">
  <div class="left">
    <h2>自動チェック</h2>
    <ul class="auto">{auto_html}</ul>
    <h2>チェック項目</h2>
    {checklist_html}
  </div>
  <div class="right">
    <h2>記事本文プレビュー</h2>
    {article_html}
  </div>
</div>
</body></html>
"""
    path.write_text(doc, encoding="utf-8")


def _extract_article_sections(article_md: str) -> List:
    """記事Markdownから (見出し, 本文プレビュー) のリストを作る。"""
    sections = []
    current = None
    buf: List[str] = []
    for line in article_md.splitlines():
        m = re.match(r"^#{2,3}\s+(.+)$", line)
        if m:
            if current:
                sections.append((current, " ".join(buf)[:200]))
            current = m.group(1).strip()
            buf = []
        elif current:
            if line.strip():
                buf.append(line.strip())
    if current:
        sections.append((current, " ".join(buf)[:200]))
    return sections


def build_review_kit(theme: str, artifacts: Dict) -> Dict:
    """監修パッケージ一式を生成する。

    artifacts には step2〜4 の出力パス・スライドオブジェクトが入る。
    """
    logger.info("=== ステップ5: 監修パッケージ生成開始（テーマ: %s）===", theme)

    article_md = _read(artifacts.get("article_md", ""))
    youtube_md = _read(artifacts.get("youtube_script", ""))
    shorts_md = _read(artifacts.get("shorts_script", ""))

    citations = CitationTracker.load(Path(artifacts.get("citations_path", "")))

    # 自動チェック
    checks = run_auto_checks(article_md, youtube_md, shorts_md)
    # Claudeリスク判定
    risk = _claude_risk_assessment(theme, article_md, checks)

    # 成果物メタ情報
    products = {
        "article_chars": count_chars(article_md),
        "youtube_chars": count_chars(youtube_md),
        "shorts_chars": count_chars(shorts_md),
        "main_slide_count": artifacts.get("main_slide_count", 0),
        "shorts_slide_count": artifacts.get("shorts_slide_count", 0),
        "main_slides": artifacts.get("main_slides", []),
        "shorts_slides": artifacts.get("shorts_slides", []),
        "article_sections": _extract_article_sections(article_md),
    }

    # 出力先
    folder = ensure_dir(OUTPUT_DIR / "review" / f"{slugify(theme)}_{today_str()}")
    xlsx_path = folder / "review_checklist.xlsx"
    md_path = folder / "review_summary.md"
    html_path = folder / "review_package.html"

    _build_excel(xlsx_path, theme, citations, checks, risk, products)
    _build_summary_md(md_path, theme, checks, risk, citations)
    _build_html_viewer(html_path, theme, checks, risk, article_md)

    logger.info("監修Excel: %s", xlsx_path)
    logger.info("監修サマリー: %s", md_path)
    logger.info("監修HTMLビューア: %s", html_path)

    return {
        "review_xlsx": str(xlsx_path),
        "review_summary": str(md_path),
        "review_html": str(html_path),
        "checks": checks,
        "risk": risk,
    }


if __name__ == "__main__":
    import sys, json

    print(json.dumps(build_review_kit(sys.argv[1], {}), ensure_ascii=False, indent=2))
