"""ステップ4b: スライド生成（python-pptx）。

YouTube本編用台本(B)とShorts用台本(C)のMarkdown表をパースして、
本編16:9とShorts 9:16のPPTXを2ファイル生成します。

デザイン:
  - 背景: settings.yaml の secondary_color
  - タイトル: 太字（本編48pt / Shorts60pt）primary_color
  - 本文: 28pt / 40pt、1スライド3行以内に丸める
  - 角丸装飾ボックス、ページ番号、ロゴ枠（placeholder）
  - 表紙＋エンドカードに監修者表記、エンドカードに免責文
  - ナレーション＋引用元IDをノート欄に挿入
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

from .utils.common import (
    DISCLAIMER_TEXT,
    OUTPUT_DIR,
    ensure_dir,
    get_logger,
    load_settings,
    reviewer_display,
    slugify,
    today_str,
)

logger = get_logger()


@dataclass
class Slide:
    """1スライド分のデータ。"""

    number: str = ""
    title: str = ""
    body: str = ""
    narration: str = ""
    ref: str = ""
    seconds: str = ""  # Shorts用の表示秒数


def _hex_to_rgb(hex_color: str) -> RGBColor:
    """'#4FB3D9' を RGBColor に変換。"""
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def parse_script_tables(markdown_text: str, is_shorts: bool = False) -> List[Slide]:
    """台本中のMarkdown表からスライド情報を抽出する。

    本編列:   枚数 | タイトル | 本文 | ナレーション | 引用元ID
    Shorts列: 枚数 | 表示秒数 | タイトル | 本文 | ナレーション | 引用元ID
    """
    slides: List[Slide] = []
    for line in markdown_text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        # 区切り行/ヘッダー行を除外
        if not cells or all(set(c) <= {"-", ":", " "} for c in cells):
            continue
        joined = "".join(cells)
        if any(h in joined for h in ["枚数", "タイトル", "ナレーション", "引用元"]) and \
                not re.search(r"\d", cells[0]):
            continue  # ヘッダー行

        try:
            if is_shorts and len(cells) >= 6:
                slides.append(Slide(
                    number=cells[0], seconds=cells[1], title=cells[2],
                    body=cells[3], narration=cells[4], ref=cells[5],
                ))
            elif len(cells) >= 5:
                slides.append(Slide(
                    number=cells[0], title=cells[1], body=cells[2],
                    narration=cells[3], ref=cells[4],
                ))
            elif len(cells) >= 2:
                # 列が足りない場合も拾えるところまで拾う
                slides.append(Slide(number=cells[0], title=cells[1],
                                    body=cells[2] if len(cells) > 2 else ""))
        except Exception as exc:
            logger.debug("表行のパース失敗: %s (%s)", line, exc)
    logger.info("台本から %d 枚のスライドを抽出", len(slides))
    return slides


def _limit_body_lines(text: str, max_lines: int = 3) -> str:
    """本文を最大3行に丸める（長文は句点等で改行）。"""
    text = re.sub(r"<br\s*/?>", "\n", text)
    parts = [p for p in re.split(r"[\n。]", text) if p.strip()]
    lines = parts[:max_lines]
    return "\n".join(s.strip().rstrip("。") + "。" for s in lines)


class SlideDeckBuilder:
    """1つのPPTXを組み立てるビルダー。"""

    def __init__(self, is_shorts: bool) -> None:
        self.is_shorts = is_shorts
        self.design = load_settings()["design"]
        self.bg = _hex_to_rgb(self.design["secondary_color"])
        self.primary = _hex_to_rgb(self.design["primary_color"])
        self.accent = _hex_to_rgb(self.design["accent_color"])
        self.title_size = Pt(60) if is_shorts else Pt(48)
        self.body_size = Pt(40) if is_shorts else Pt(28)

        self.prs = Presentation()
        if is_shorts:  # 9:16
            self.prs.slide_width = Emu(6858000)   # 7.5 in
            self.prs.slide_height = Emu(12192000)  # 13.333 in
        else:  # 16:9
            self.prs.slide_width = Emu(12192000)
            self.prs.slide_height = Emu(6858000)
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height

    # -------------------------------------------------------------- helpers
    def _blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        # 背景色
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.bg
        return slide

    def _add_textbox(self, slide, left, top, width, height, text, size,
                     color, bold=False, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        first = True
        for ln in str(text).split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            run = p.add_run()
            run.text = ln
            run.font.size = size
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = self.design["font_main"]
        return box

    def _rounded_box(self, slide, left, top, width, height, color):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = color
        shape.shadow.inherit = False
        return shape

    def _logo_placeholder(self, slide):
        """右上にロゴ枠（差替可能なプレースホルダー）。"""
        size = Emu(700000)
        box = self._rounded_box(slide, self.W - size - Emu(200000), Emu(150000),
                                 size, Emu(350000), self.primary)
        tf = box.text_frame
        tf.text = "LOGO"
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].runs[0]
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(255, 255, 255)

    def _page_number(self, slide, n):
        box = self._add_textbox(
            slide, self.W - Emu(900000), self.H - Emu(500000),
            Emu(700000), Emu(350000), str(n), Pt(14),
            RGBColor(150, 150, 150), align=PP_ALIGN.RIGHT,
        )
        return box

    def _set_notes(self, slide, narration, ref):
        note = (narration or "").strip()
        if ref:
            note += f"\n\n[引用元ID: {ref}]"
        slide.notes_slide.notes_text_frame.text = note

    # ---------------------------------------------------------------- slides
    def add_cover(self, title: str):
        slide = self._blank()
        self._logo_placeholder(slide)
        margin = Emu(600000)
        # タイトル装飾ボックス
        self._rounded_box(slide, margin, self.H // 3, self.W - 2 * margin,
                          self.H // 4, self.accent)
        self._add_textbox(
            slide, margin + Emu(200000), self.H // 3 + Emu(150000),
            self.W - 2 * margin - Emu(400000), self.H // 4 - Emu(200000),
            title, self.title_size, RGBColor(255, 255, 255), bold=True,
            align=PP_ALIGN.CENTER,
        )
        # 監修者表記（表紙）
        self._add_textbox(
            slide, margin, self.H - Emu(900000), self.W - 2 * margin,
            Emu(500000), reviewer_display(), Pt(24) if not self.is_shorts else Pt(30),
            self.primary, bold=True, align=PP_ALIGN.CENTER,
        )

    def add_content(self, s: Slide, page: int):
        slide = self._blank()
        self._logo_placeholder(slide)
        margin = Emu(500000)
        # タイトル
        self._add_textbox(
            slide, margin, Emu(600000), self.W - 2 * margin, Emu(1200000),
            s.title, self.title_size, self.primary, bold=True,
        )
        # 本文（角丸ボックス内、最大3行）
        body_top = Emu(2200000) if not self.is_shorts else Emu(3000000)
        self._rounded_box(slide, margin, body_top, self.W - 2 * margin,
                          self.H // 3, RGBColor(255, 255, 255))
        self._add_textbox(
            slide, margin + Emu(200000), body_top + Emu(150000),
            self.W - 2 * margin - Emu(400000), self.H // 3 - Emu(200000),
            _limit_body_lines(s.body), self.body_size, RGBColor(60, 60, 60),
        )
        if self.is_shorts and s.seconds:
            self._add_textbox(
                slide, margin, Emu(300000), Emu(2000000), Emu(400000),
                f"⏱ {s.seconds}", Pt(28), self.accent, bold=True,
            )
        self._page_number(slide, page)
        self._set_notes(slide, s.narration, s.ref)

    def add_endcard(self):
        slide = self._blank()
        self._logo_placeholder(slide)
        margin = Emu(600000)
        self._add_textbox(
            slide, margin, self.H // 4, self.W - 2 * margin, Emu(800000),
            "まとめ", self.title_size, self.primary, bold=True,
            align=PP_ALIGN.CENTER,
        )
        # 監修者表記（エンドカード）
        self._add_textbox(
            slide, margin, self.H // 2, self.W - 2 * margin, Emu(600000),
            reviewer_display(), Pt(28) if not self.is_shorts else Pt(34),
            self.primary, bold=True, align=PP_ALIGN.CENTER,
        )
        # 免責文
        self._add_textbox(
            slide, margin, self.H // 2 + Emu(800000), self.W - 2 * margin,
            Emu(1500000), DISCLAIMER_TEXT, Pt(18) if not self.is_shorts else Pt(24),
            RGBColor(110, 110, 110), align=PP_ALIGN.CENTER,
        )

    def save(self, path: Path) -> Path:
        self.prs.save(str(path))
        return path


def _build_one(title: str, slides: List[Slide], is_shorts: bool, path: Path) -> Path:
    builder = SlideDeckBuilder(is_shorts=is_shorts)
    builder.add_cover(title)
    page = 1
    for s in slides:
        builder.add_content(s, page)
        page += 1
    builder.add_endcard()
    builder.save(path)
    return path


def build_slides(theme: str, youtube_script_path: str, shorts_script_path: str) -> dict:
    """本編・ShortsのPPTXを生成して保存する。"""
    logger.info("=== ステップ4b: スライド生成開始（テーマ: %s）===", theme)

    yt_md = Path(youtube_script_path).read_text(encoding="utf-8")
    sh_md = Path(shorts_script_path).read_text(encoding="utf-8")

    yt_slides = parse_script_tables(yt_md, is_shorts=False)
    sh_slides = parse_script_tables(sh_md, is_shorts=True)

    # 表が空でも最低限のスライドを作る（フォールバック）
    if not yt_slides:
        yt_slides = [Slide(title=theme, body="台本表の抽出に失敗しました。台本ファイルをご確認ください。")]
    if not sh_slides:
        sh_slides = [Slide(title=theme, body="台本表の抽出に失敗しました。", seconds="5秒")]

    slug = slugify(theme)
    base = ensure_dir(OUTPUT_DIR / "slides")
    main_path = base / f"{slug}_{today_str()}_main.pptx"
    shorts_path = base / f"{slug}_{today_str()}_shorts.pptx"

    _build_one(theme, yt_slides, False, main_path)
    _build_one(theme, sh_slides, True, shorts_path)

    logger.info("本編スライド: %s（%d枚）", main_path, len(yt_slides))
    logger.info("Shortsスライド: %s（%d枚）", shorts_path, len(sh_slides))

    return {
        "slides_main": str(main_path),
        "slides_shorts": str(shorts_path),
        "main_slide_count": len(yt_slides),
        "shorts_slide_count": len(sh_slides),
        "main_slides": yt_slides,
        "shorts_slides": sh_slides,
    }


if __name__ == "__main__":
    import sys

    print(build_slides(sys.argv[1], sys.argv[2], sys.argv[3]))
