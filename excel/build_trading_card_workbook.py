# -*- coding: utf-8 -*-
"""トレカ事業 従業員管理表(Excel) 生成スクリプト.

レイアウト:
  「月別入力」シートに、1か月＝1ブロックを縦に12か月分並べる。
  各ブロックの中では従業員ごとに独立した表を左右に並列配置する。
    ・左(A〜H列)  = 従業員A の表
    ・右(J〜Q列)  = 従業員B の表
  各表の項目: 日付 / 仕入れ店舗 / 仕入れ合計金額 / 支払方法 / 立替(○) /
              移動区間(出発) / 移動区間(到着) / 交通費
  移動区間が複数ある日は、表の途中に行を挿入して増やせる(小計は自動拡張)。
  ブロック下部に共通欄(送料・その他経費 内容/金額)と当月サマリーを置く。
  「年間集計」シートで12か月分を横断集計する。
"""

from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

FONT = "Meiryo"

# ---- 配色 -------------------------------------------------------------------
C_TITLE = "1F3864"     # 濃紺
C_MONTH = "1F3864"     # 月見出し
C_EMP_A = "2E5C8A"     # 従業員A 見出し(青)
C_EMP_B = "4A7A3F"     # 従業員B 見出し(緑)
C_FEE = "8A5A2E"       # 共通経費 見出し(茶)
C_SUM = "5B3A7A"       # サマリー 見出し(紫)
C_INPUT = "FFF7DC"     # 入力セル(淡黄)
C_AUTO = "EDEDED"      # 自動計算セル(グレー)
C_SAMPLE = "E8F1FB"    # 記入例
C_SUBA = "DCE6F1"      # 小計行(A)
C_SUBB = "E2EFDA"      # 小計行(B)
C_SUBF = "FDF0D5"      # 小計行(共通経費)
C_TOTAL = "FFF2CC"     # 当月総支出

THIN = Side(style="thin", color="B7C3CE")
MED = Side(style="medium", color="2E5C8A")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_HEAD = Border(left=THIN, right=THIN, top=MED, bottom=MED)

FILL_INPUT = PatternFill("solid", fgColor=C_INPUT)
FILL_AUTO = PatternFill("solid", fgColor=C_AUTO)
FILL_SAMPLE = PatternFill("solid", fgColor=C_SAMPLE)

F_TITLE = Font(name=FONT, size=14, bold=True, color=C_TITLE)
F_NOTE = Font(name=FONT, size=9, color="555555")
F_MONTH = Font(name=FONT, size=12, bold=True, color="FFFFFF")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")     # 手入力は青字
F_AUTO = Font(name=FONT, size=10, color="404040")      # 自動計算
F_SAMPLE = Font(name=FONT, size=10, color="7F7F7F", italic=True)
F_SUB = Font(name=FONT, size=10, bold=True)
F_GRAND = Font(name=FONT, size=11, bold=True, color=C_TITLE)

YEN = '#,##0;-#,##0;"-"'
DATE_FMT = "yyyy/mm/dd"

MONTHS = 12          # 何か月分のブロックを作るか
N_ROWS = 25          # 従業員1人あたり1か月の入力行数
M_ROWS = 8           # 共通経費の入力行数

# ---- ブロック内のオフセット --------------------------------------------------
O_MONTH = 0                       # 月見出し
O_EMP = 1                         # 従業員名見出し
O_HEAD = 2                        # 列ヘッダー
O_FIRST = 3                       # 入力開始
O_LAST = O_FIRST + N_ROWS - 1     # 入力終了 (=27)
O_SUB = O_LAST + 1                # 小計 (=28)
O_SEC = O_SUB + 2                 # 「共通経費」「当月サマリー」見出し (=30)
O_SECHEAD = O_SEC + 1             # その列ヘッダー (=31)
O_FEE_FIRST = O_SECHEAD + 1       # 共通経費 入力開始 (=32)
O_FEE_LAST = O_FEE_FIRST + M_ROWS - 1   # (=39)
O_FEE_SUB = O_FEE_LAST + 1        # 共通経費 小計 / 当月総支出 (=40)
BLOCK_H = O_FEE_SUB + 3           # 次ブロックまで (=43)
FIRST_BLOCK_ROW = 7

# 従業員A: A〜H列 / 区切り: I列 / 従業員B: J〜Q列
A_COL0 = 1
B_COL0 = 10
COLS = [
    ("日付", 11),
    ("仕入れ店舗", 20),
    ("仕入れ合計金額（円）", 16),
    ("支払方法", 14),
    ("立替\n(○)", 6),
    ("移動区間\n(出発)", 14),
    ("移動区間\n(到着)", 14),
    ("交通費（円）", 12),
]
C_DATE, C_SHOP, C_AMT, C_PAY, C_ADV, C_FROM, C_TO, C_FARE = range(8)

PAY_METHODS = ["現金", "クレジットカード", "電子マネー", "QRコード決済", "銀行振込", "その他"]
FEE_KINDS = ["送料", "その他経費"]
START_MONTH = date(2026, 8, 1)


def box(ws, r1, c1, r2, c2, border=BORDER, fill=None, font=None, align=None):
    """矩形範囲に罫線・塗り・フォントをまとめて適用する(結合セル対応)。"""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font
            if align is not None:
                cell.alignment = align


wb = Workbook()

# ============================================================================
# 設定(マスタ)シート
# ============================================================================
st = wb.active
st.title = "設定"
st["A1"] = "設定(マスタ)"
st["A1"].font = F_TITLE
st["A2"] = "ここを書き換えると、各シートの見出しやプルダウンに反映されます。"
st["A2"].font = F_NOTE

for coord, label in (("B4", "従業員名"), ("D4", "支払方法"), ("F4", "経費区分"), ("H4", "立替")):
    st[coord] = label
    st[coord].font = F_HEAD
    st[coord].fill = PatternFill("solid", fgColor=C_EMP_A)
    st[coord].alignment = Alignment(horizontal="center", vertical="center")
    st[coord].border = BORDER_HEAD

def put_list(col, values):
    for i, v in enumerate(values):
        cell = st.cell(row=5 + i, column=col, value=v)
        cell.font = F_INPUT
        cell.fill = FILL_INPUT
        cell.border = BORDER

put_list(2, ["従業員A", "従業員B"])
put_list(4, PAY_METHODS)
put_list(6, FEE_KINDS)
put_list(8, ["○"])
st["H5"].alignment = Alignment(horizontal="center")

st["A12"] = "開始年月"
st["A12"].font = F_SUB
st["B12"] = START_MONTH
st["B12"].number_format = "yyyy/mm"
st["B12"].font = F_INPUT
st["B12"].fill = FILL_INPUT
st["B12"].border = BORDER
st["B12"].alignment = Alignment(horizontal="center")
st["B12"].comment = Comment(
    "「月別入力」シートの1ブロック目の年月です。\n"
    "この日付を変えると、12か月分すべての月見出しが自動でずれます。\n"
    "必ず「その月の1日」を入力してください（例：2026/08/01）。", "管理表")
st["C12"] = "← 月別入力シートの1ブロック目の年月（毎月1日を入力）"
st["C12"].font = F_NOTE
st["A13"] = "※従業員名を変更すると、月別入力・年間集計の見出しも自動で変わります。"
st["A13"].font = F_NOTE
st["A14"] = "※従業員が3人以上になる場合は、表の追加が必要です（このスクリプトを修正してください）。"
st["A14"].font = F_NOTE

for col, w in (("A", 12), ("B", 18), ("C", 46), ("D", 18), ("E", 3), ("F", 16), ("G", 3), ("H", 8)):
    st.column_dimensions[col].width = w

EMP_REF = ["設定!$B$5", "設定!$B$6"]

# ============================================================================
# 月別入力シート
# ============================================================================
ws = wb.create_sheet("月別入力")
ws.sheet_view.showGridLines = False

ws["A1"] = "トレカ事業　従業員 仕入れ・交通費 管理表（月別／従業員別）"
ws["A1"].font = F_TITLE
ws["A2"] = "【凡例】黄色セル＝手入力（青字） ／ グレー・色付きセル＝自動計算（編集しないでください）"
ws["A2"].font = F_NOTE
ws["A3"] = ("【構成】1か月＝1ブロック。ブロック内の左（A〜H列）が従業員A、右（J〜Q列）が従業員Bの独立した表です。"
            "ブロックは上から12か月分並んでいます。")
ws["A3"].font = F_NOTE
ws["A4"] = ("【交通費】1行に1区間を入力します。同じ日に移動区間が複数ある場合は行を増やし、"
            "日付を再入力して区間と交通費だけを記入してください（仕入れ金額はその日の1行目のみ）。"
            "行が足りないときは表の途中（最終行の1つ上まで）で行を挿入すると、小計の範囲も自動で広がります。")
ws["A4"].font = F_NOTE
ws["A5"] = ("【立替】従業員が自分で支払った行に「○」を入力してください。"
            "その行の仕入れ合計金額＋交通費が「立替精算額（返金額）」に集計されます。")
ws["A5"].font = F_NOTE

for i, (_, width) in enumerate(COLS):
    ws.column_dimensions[get_column_letter(A_COL0 + i)].width = width
    ws.column_dimensions[get_column_letter(B_COL0 + i)].width = width
ws.column_dimensions[get_column_letter(B_COL0 - 1)].width = 2   # I列(区切り)

LAST_COL = B_COL0 + len(COLS) - 1   # Q列

dv_pay = DataValidation(type="list", formula1="=設定!$D$5:$D$10", allow_blank=True,
                        showDropDown=False, promptTitle="支払方法",
                        prompt="現金／クレジットカード等を選択してください。")
dv_adv = DataValidation(type="list", formula1='"○"', allow_blank=True,
                        showDropDown=False, promptTitle="立替",
                        prompt="従業員が立て替えた行のみ ○ を入力してください。")
dv_kind = DataValidation(type="list", formula1="=設定!$F$5:$F$6", allow_blank=True,
                         showDropDown=False, promptTitle="経費区分",
                         prompt="送料／その他経費 を選択してください。")
for dv in (dv_pay, dv_adv, dv_kind):
    ws.add_data_validation(dv)

blocks = []   # 年間集計から参照するための行番号を記録

for m in range(MONTHS):
    R = FIRST_BLOCK_ROW + m * BLOCK_H
    r_month = R + O_MONTH
    r_emp = R + O_EMP
    r_head = R + O_HEAD
    r_first = R + O_FIRST
    r_last = R + O_LAST
    r_sub = R + O_SUB
    r_sec = R + O_SEC
    r_sechead = R + O_SECHEAD
    r_fee_first = R + O_FEE_FIRST
    r_fee_last = R + O_FEE_LAST
    r_fee_sub = R + O_FEE_SUB

    # --- 月見出し -----------------------------------------------------------
    ws.merge_cells(start_row=r_month, start_column=1, end_row=r_month, end_column=LAST_COL)
    box(ws, r_month, 1, r_month, LAST_COL,
        border=Border(left=MED, right=MED, top=MED, bottom=MED),
        fill=PatternFill("solid", fgColor=C_MONTH), font=F_MONTH,
        align=Alignment(horizontal="left", vertical="center", indent=1))
    ws.cell(row=r_month, column=1,
            value=f'=TEXT(EDATE(設定!$B$12,{m}),"yyyy年m月")')
    ws.row_dimensions[r_month].height = 24

    # --- 従業員ごとの独立した表(左右に並列) ---------------------------------
    for e, (col0, color, sub_fill) in enumerate(
            ((A_COL0, C_EMP_A, C_SUBA), (B_COL0, C_EMP_B, C_SUBB))):
        c_end = col0 + len(COLS) - 1

        ws.merge_cells(start_row=r_emp, start_column=col0, end_row=r_emp, end_column=c_end)
        box(ws, r_emp, col0, r_emp, c_end, border=BORDER_HEAD,
            fill=PatternFill("solid", fgColor=color), font=F_HEAD,
            align=Alignment(horizontal="center", vertical="center"))
        ws.cell(row=r_emp, column=col0, value=f"={EMP_REF[e]}")

        for i, (title, _) in enumerate(COLS):
            c = ws.cell(row=r_head, column=col0 + i, value=title)
            c.font = F_HEAD
            c.fill = PatternFill("solid", fgColor=color)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER_HEAD
        ws.row_dimensions[r_head].height = 32

        # 入力行
        box(ws, r_first, col0, r_last, c_end, fill=FILL_INPUT, font=F_INPUT)
        for r in range(r_first, r_last + 1):
            ws.cell(row=r, column=col0 + C_DATE).number_format = DATE_FMT
            ws.cell(row=r, column=col0 + C_AMT).number_format = YEN
            ws.cell(row=r, column=col0 + C_FARE).number_format = YEN
            ws.cell(row=r, column=col0 + C_ADV).alignment = Alignment(horizontal="center")

        dv_pay.add(f"{get_column_letter(col0 + C_PAY)}{r_first}:"
                   f"{get_column_letter(col0 + C_PAY)}{r_last}")
        dv_adv.add(f"{get_column_letter(col0 + C_ADV)}{r_first}:"
                   f"{get_column_letter(col0 + C_ADV)}{r_last}")

        # 小計行
        box(ws, r_sub, col0, r_sub, c_end, border=BORDER_HEAD,
            fill=PatternFill("solid", fgColor=sub_fill), font=F_SUB)
        ws.merge_cells(start_row=r_sub, start_column=col0, end_row=r_sub, end_column=col0 + 1)
        ws.cell(row=r_sub, column=col0, value="小計").alignment = Alignment(
            horizontal="center", vertical="center")
        amt = get_column_letter(col0 + C_AMT)
        fare = get_column_letter(col0 + C_FARE)
        c = ws.cell(row=r_sub, column=col0 + C_AMT, value=f"=SUM({amt}{r_first}:{amt}{r_last})")
        c.number_format = YEN
        c = ws.cell(row=r_sub, column=col0 + C_FARE, value=f"=SUM({fare}{r_first}:{fare}{r_last})")
        c.number_format = YEN

    # --- 共通経費(送料・その他経費) -----------------------------------------
    ws.merge_cells(start_row=r_sec, start_column=1, end_row=r_sec, end_column=8)
    box(ws, r_sec, 1, r_sec, 8, border=BORDER_HEAD,
        fill=PatternFill("solid", fgColor=C_FEE), font=F_HEAD,
        align=Alignment(horizontal="center", vertical="center"))
    ws.cell(row=r_sec, column=1, value="共通経費（送料・その他経費）　※従業員個人に紐づかない経費")

    fee_cols = [("日付", 1, 1), ("区分", 2, 2), ("内容", 3, 6), ("金額（円）", 7, 8)]
    for title, c1, c2 in fee_cols:
        if c1 != c2:
            ws.merge_cells(start_row=r_sechead, start_column=c1, end_row=r_sechead, end_column=c2)
        box(ws, r_sechead, c1, r_sechead, c2, border=BORDER_HEAD,
            fill=PatternFill("solid", fgColor=C_FEE), font=F_HEAD,
            align=Alignment(horizontal="center", vertical="center"))
        ws.cell(row=r_sechead, column=c1, value=title)

    for r in range(r_fee_first, r_fee_last + 1):
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        box(ws, r, 1, r, 8, fill=FILL_INPUT, font=F_INPUT)
        ws.cell(row=r, column=1).number_format = DATE_FMT
        ws.cell(row=r, column=7).number_format = YEN
    dv_kind.add(f"B{r_fee_first}:B{r_fee_last}")

    ws.merge_cells(start_row=r_fee_sub, start_column=1, end_row=r_fee_sub, end_column=6)
    ws.merge_cells(start_row=r_fee_sub, start_column=7, end_row=r_fee_sub, end_column=8)
    box(ws, r_fee_sub, 1, r_fee_sub, 8, border=BORDER_HEAD,
        fill=PatternFill("solid", fgColor=C_SUBF), font=F_SUB)
    ws.cell(row=r_fee_sub, column=1, value="共通経費 小計").alignment = Alignment(
        horizontal="center", vertical="center")
    c = ws.cell(row=r_fee_sub, column=7, value=f"=SUM(G{r_fee_first}:G{r_fee_last})")
    c.number_format = YEN

    # --- 当月サマリー(右側 J〜M列) ------------------------------------------
    ws.merge_cells(start_row=r_sec, start_column=B_COL0, end_row=r_sec, end_column=B_COL0 + 3)
    box(ws, r_sec, B_COL0, r_sec, B_COL0 + 3, border=BORDER_HEAD,
        fill=PatternFill("solid", fgColor=C_SUM), font=F_HEAD,
        align=Alignment(horizontal="center", vertical="center"))
    ws.cell(row=r_sec, column=B_COL0, value="当月サマリー")

    ws.merge_cells(start_row=r_sechead, start_column=B_COL0, end_row=r_sechead, end_column=B_COL0 + 2)
    box(ws, r_sechead, B_COL0, r_sechead, B_COL0 + 3, border=BORDER_HEAD,
        fill=PatternFill("solid", fgColor=C_SUM), font=F_HEAD,
        align=Alignment(horizontal="center", vertical="center"))
    ws.cell(row=r_sechead, column=B_COL0, value="項目")
    ws.cell(row=r_sechead, column=B_COL0 + 3, value="金額（円）")

    a0, b0 = A_COL0, B_COL0
    def emp_range(col0, idx):
        col = get_column_letter(col0 + idx)
        return f"{col}{r_first}:{col}{r_last}"

    def adv_formula(col0):
        adv = emp_range(col0, C_ADV)
        return (f'=SUMIF({adv},"○",{emp_range(col0, C_AMT)})'
                f'+SUMIF({adv},"○",{emp_range(col0, C_FARE)})')

    summary = [
        (f'={EMP_REF[0]}&" 仕入れ合計"', f"={get_column_letter(a0 + C_AMT)}{r_sub}"),
        (f'={EMP_REF[0]}&" 交通費合計"', f"={get_column_letter(a0 + C_FARE)}{r_sub}"),
        (f'={EMP_REF[0]}&" 立替精算額（返金額）"', adv_formula(a0)),
        (f'={EMP_REF[1]}&" 仕入れ合計"', f"={get_column_letter(b0 + C_AMT)}{r_sub}"),
        (f'={EMP_REF[1]}&" 交通費合計"', f"={get_column_letter(b0 + C_FARE)}{r_sub}"),
        (f'={EMP_REF[1]}&" 立替精算額（返金額）"', adv_formula(b0)),
        ('="共通経費　送料"', f'=SUMIF(B{r_fee_first}:B{r_fee_last},"送料",G{r_fee_first}:G{r_fee_last})'),
        ('="共通経費　その他経費"',
         f'=SUMIF(B{r_fee_first}:B{r_fee_last},"その他経費",G{r_fee_first}:G{r_fee_last})'),
    ]
    sum_rows = {}
    for i, (label, formula) in enumerate(summary):
        r = r_fee_first + i
        ws.merge_cells(start_row=r, start_column=b0, end_row=r, end_column=b0 + 2)
        box(ws, r, b0, r, b0 + 3, fill=FILL_AUTO, font=F_AUTO)
        ws.cell(row=r, column=b0, value=label).alignment = Alignment(vertical="center", indent=1)
        c = ws.cell(row=r, column=b0 + 3, value=formula)
        c.number_format = YEN
        sum_rows[i] = r

    # 当月総支出
    ws.merge_cells(start_row=r_fee_sub, start_column=b0, end_row=r_fee_sub, end_column=b0 + 2)
    box(ws, r_fee_sub, b0, r_fee_sub, b0 + 3,
        border=Border(left=MED, right=MED, top=MED, bottom=MED),
        fill=PatternFill("solid", fgColor=C_TOTAL), font=F_GRAND)
    ws.cell(row=r_fee_sub, column=b0, value="当月 総支出").alignment = Alignment(
        horizontal="center", vertical="center")
    mcol = get_column_letter(b0 + 3)
    c = ws.cell(row=r_fee_sub, column=b0 + 3,
                value=f"={mcol}{sum_rows[0]}+{mcol}{sum_rows[1]}+{mcol}{sum_rows[3]}"
                      f"+{mcol}{sum_rows[4]}+G{r_fee_sub}")
    c.number_format = YEN

    blocks.append({
        "month_row": r_month, "sum_col": mcol, "sum_rows": sum_rows,
        "fee_sub_row": r_fee_sub,
    })

    if m > 0:
        ws.row_breaks.append(Break(id=r_month - 1))

ws.freeze_panes = "A6"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ============================================================================
# 年間集計シート
# ============================================================================
sm = wb.create_sheet("年間集計")
sm.sheet_view.showGridLines = False
sm["A1"] = "年間集計（12か月）"
sm["A1"].font = F_TITLE
sm["A2"] = "「月別入力」シートの各月ブロックのサマリーを、そのまま参照しています（入力欄はありません）。"
sm["A2"].font = F_NOTE

head = ["年月", "A 仕入れ", "A 交通費", "A 立替精算額", "B 仕入れ", "B 交通費",
        "B 立替精算額", "共通 送料", "共通 その他経費", "当月 総支出"]
HR = 4
for i, t in enumerate(head, start=1):
    c = sm.cell(row=HR, column=i, value=t)
    c.font = F_HEAD
    c.fill = PatternFill("solid", fgColor=C_SUM)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER_HEAD
sm.row_dimensions[HR].height = 32
sm.cell(row=HR, column=2).comment = Comment(
    "見出しの A / B は「設定」シートの従業員名に対応します（A＝上、B＝下）。", "管理表")

IN = "月別入力"
for m, blk in enumerate(blocks):
    r = HR + 1 + m
    sm.cell(row=r, column=1, value=f"={IN}!A{blk['month_row']}")
    col = blk["sum_col"]
    for i in range(8):
        sm.cell(row=r, column=2 + i, value=f"={IN}!{col}{blk['sum_rows'][i]}")
    sm.cell(row=r, column=10, value=f"={IN}!{col}{blk['fee_sub_row']}")
    box(sm, r, 1, r, 10, fill=FILL_AUTO, font=F_AUTO)
    for c in range(2, 11):
        sm.cell(row=r, column=c).number_format = YEN
    sm.cell(row=r, column=1).alignment = Alignment(horizontal="center")

TR = HR + 1 + MONTHS
sm.cell(row=TR, column=1, value="年間合計").alignment = Alignment(horizontal="center")
for c in range(2, 11):
    col = get_column_letter(c)
    cell = sm.cell(row=TR, column=c, value=f"=SUM({col}{HR+1}:{col}{TR-1})")
    cell.number_format = YEN
box(sm, TR, 1, TR, 10, border=Border(left=MED, right=MED, top=MED, bottom=MED),
    fill=PatternFill("solid", fgColor=C_TOTAL), font=F_GRAND)

sm.cell(row=TR + 2, column=1,
        value="※「立替精算額（返金額）」＝立替欄に○が付いた行の 仕入れ合計金額＋交通費。従業員へ返金する金額です。").font = F_NOTE
sm.cell(row=TR + 3, column=1,
        value="※「当月 総支出」＝両名の 仕入れ＋交通費 ＋ 共通経費（送料・その他経費）。支払方法や立替の有無は問いません。").font = F_NOTE

sm.column_dimensions["A"].width = 14
for c in range(2, 11):
    sm.column_dimensions[get_column_letter(c)].width = 15
sm.freeze_panes = "B5"

# ============================================================================
# 記入例(1ブロック目のみ)
# ============================================================================
b0_row = FIRST_BLOCK_ROW + O_FIRST
samples_a = [
    (date(2026, 8, 1), "カードショップ〇〇 秋葉原店", 45000, "現金", "○", "自宅", "秋葉原", 480),
    (date(2026, 8, 1), None, None, None, "○", "秋葉原", "中野", 200),
]
samples_b = [
    (date(2026, 8, 2), "リサイクルショップ△△ 大宮店", 128000, "クレジットカード", None, "自宅", "大宮", 640),
]
for col0, rows in ((A_COL0, samples_a), (B_COL0, samples_b)):
    for i, row in enumerate(rows):
        r = b0_row + i
        for j, val in enumerate(row):
            cell = ws.cell(row=r, column=col0 + j, value=val)
            cell.font = F_SAMPLE
            cell.fill = FILL_SAMPLE
        ws.cell(row=r, column=col0 + C_DATE).number_format = DATE_FMT
        ws.cell(row=r, column=col0 + C_AMT).number_format = YEN
        ws.cell(row=r, column=col0 + C_FARE).number_format = YEN
        ws.cell(row=r, column=col0 + C_ADV).alignment = Alignment(horizontal="center")

fee_row0 = FIRST_BLOCK_ROW + O_FEE_FIRST
fee_samples = [
    (date(2026, 8, 3), "送料", "メルカリ発送（ゆうパケットプラス）×5件", 2600),
    (date(2026, 8, 5), "その他経費", "スリーブ・ローダー購入", 3480),
]
for i, (d, kind, memo, amount) in enumerate(fee_samples):
    r = fee_row0 + i
    for c, val in ((1, d), (2, kind), (3, memo), (7, amount)):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = F_SAMPLE
        cell.fill = FILL_SAMPLE
    box(ws, r, 1, r, 8, fill=FILL_SAMPLE, font=F_SAMPLE)
    ws.cell(row=r, column=1).number_format = DATE_FMT
    ws.cell(row=r, column=7).number_format = YEN

ws.cell(row=FIRST_BLOCK_ROW + O_SEC - 1, column=1,
        value="↑ 上の色付きの行は記入例です。ご利用前に内容を削除してください。").font = F_NOTE

# ---- フォントを全セルに適用 --------------------------------------------------
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            f = cell.font
            if f.name != FONT:
                cell.font = Font(name=FONT, size=f.size or 10, bold=f.bold,
                                 italic=f.italic, color=f.color)

wb.active = wb.sheetnames.index("月別入力")
OUT = "/home/user/toto/excel/トレカ事業_従業員管理表.xlsx"
wb.save(OUT)
print("saved:", OUT)


def fix_fonts_in_styles(path=OUT, font=FONT):
    """recalc.py(LibreOffice)で再計算した後に実行する後処理.

    LibreOffice は Meiryo を持たない環境で styles.xml のフォント名を代替フォントに
    書き換えてしまう。数式のキャッシュ値を壊さないよう xl/styles.xml だけを直接置換する。
    """
    import os
    import zipfile

    tmp = path + ".tmp"
    replaced = 0
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                text = data.decode("utf-8")
                for alt in ("WenQuanYi Zen Hei", "Liberation Sans", "DejaVu Sans"):
                    replaced += text.count(alt)
                    text = text.replace(alt, font)
                # 既定フォント(Calibri)も置換し、空白セルに入力しても書体が揃うようにする
                replaced += text.count('val="Calibri"')
                text = text.replace('val="Calibri"', f'val="{font}"')
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)
    return replaced
